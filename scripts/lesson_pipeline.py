"""Deterministic gates for the agent-authored lesson and local media pipeline.

The agent creates/reviews content and images. This CLI validates, records approval,
resumes unchanged stages and rejects invalid media instead of claiming success.
"""
import argparse
import hashlib
import json
import math
import os
from pathlib import Path
import shutil
import subprocess
import sys
import time
import unicodedata

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')

from doc_reader import scan_root_directory
from parse_srt import parse_srt

SCRIPTS = Path(__file__).resolve().parent
STOP = 'Đã hoàn thành đến bước 4, cần bạn đánh giá chốt nội dung.'


def read(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def save(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + '.tmp')
    temporary.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    temporary.replace(path)


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def file_hash(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def require(condition, message):
    if not condition:
        raise ValueError(message)


def validate_plan(plan):
    for key in ('title', 'assessment', 'changes', 'objectives', 'slides'):
        require(plan.get(key), f'Thiếu {key}')
    for level in ('know', 'understand', 'apply'):
        require(plan['objectives'].get(level), f'Thiếu mục tiêu {level}')
    slides = plan['slides']
    require(10 <= len(slides) <= 13, 'Cần 10–13 slide')
    require([s['id'] for s in slides] == list(range(1, len(slides) + 1)), 'STT phải liên tục từ 1')
    require(slides[0]['role'] == 'introduction' and slides[1]['role'] == 'objectives',
            'Slide 1 giới thiệu, slide 2 mục tiêu')
    require({'foundation', 'content', 'example', 'practice', 'summary'} <= {s['role'] for s in slides},
            'Thiếu kiến thức nền/nội dung/ví dụ/thực hành/tổng kết')
    for slide in slides:
        require(slide.get('title') and slide.get('main_idea') and slide.get('source_refs'),
                f"Slide {slide['id']}: cần tên, một ý chính và nguồn")


def validate_storyboard(board, plan):
    slides = board['slides']
    require(len(slides) == len(plan['slides']), 'Kịch bản khác số slide đã duyệt')
    for slide, outline in zip(slides, plan['slides']):
        require(slide['id'] == outline['id'] and slide['main_idea'] == outline['main_idea'],
                'Kịch bản thay đổi ý chính đã duyệt; quay lại bước 4')
        bullets = slide['bullets']
        require(3 <= len(bullets) <= 7, f"Slide {slide['id']}: cần 3–7 ý ngắn")
        require(all(isinstance(b, str) and 0 < len(b.split()) <= 15 for b in bullets),
                f"Slide {slide['id']}: ý vượt 15 từ hoặc rỗng")
        require(60 <= len(slide['voice_script'].split()) <= 120,
                f"Slide {slide['id']}: lời giảng phải 60–120 từ")
        prompt = slide['image_prompt'].lower()
        require(all(term in prompt for term in ('16:9', 'nền trắng', 'bên trái', 'bên phải', 'be vietnam pro')),
                f"Slide {slide['id']}: prompt thiếu quy chuẩn hình ảnh")


def sources_current(run):
    old = read(run / 'sources.json')
    current = scan_root_directory(old['root'])
    require(not current['errors'] and current['documents'], 'Tài liệu nguồn thiếu hoặc không đọc được')
    signature = lambda data: [(d['path'], d['sha256']) for d in data['documents']]
    require(signature(old) == signature(current), 'Nguồn đã đổi: chạy init với thư mục output mới và duyệt lại')


def approved(run):
    state = read(run / 'state.json')
    sources_current(run)
    plan = read(run / 'plan.json')
    validate_plan(plan)
    require(state.get('approved_plan') == digest(plan), 'Chưa duyệt hoặc khung bài đã đổi: phải dừng sau bước 4')
    return state, plan


def command(args):
    subprocess.run([str(a) for a in args], check=True, env={**os.environ, 'PYTHONUTF8': '1'})


def probe(path):
    require(shutil.which('ffprobe'), 'Thiếu ffprobe trong PATH')
    result = subprocess.run(['ffprobe', '-v', 'error', '-show_streams', '-show_format',
                             '-of', 'json', str(path)], capture_output=True, text=True, check=True)
    return json.loads(result.stdout)


def duration(path):
    seconds = float(probe(path)['format']['duration'])
    require(math.isfinite(seconds) and seconds > 0, f'Thời lượng không hợp lệ: {path}')
    return seconds


def validate_srt(path, duration_ms):
    text = Path(path).read_text(encoding='utf-8-sig')
    cues = parse_srt(text)
    require(cues and text.count('-->') == len(cues), f'SRT rỗng/sai cấu trúc: {path}')
    previous = 0
    for cue in cues:
        require(previous <= cue['startMs'] < cue['endMs'] <= duration_ms + 100 and cue['text'],
                f'SRT sai thứ tự hoặc vượt audio: {path}')
        previous = cue['endMs']
    return cues


def validate_annotation(path, size, total_ms):
    annotation = read(path)
    require(annotation['canvas'] == {'width': size[0], 'height': size[1]}, 'Canvas không khớp ảnh')
    elements = annotation['elements']
    require(elements, 'Annotation rỗng')
    sequences = [e['sequence'] for e in elements]
    require(sequences == sorted(set(sequences)), 'Sequence phải duy nhất và tăng dần')
    previous = 0
    for element in elements:
        region = element['region']
        reveal = element['reveal']
        for rect in [region] + reveal.get('protectedRegions', []):
            x, y, w, h = [rect[k] for k in ('x', 'y', 'width', 'height')]
            require(all(type(v) is int for v in (x, y, w, h)) and min(x, y) >= 0
                    and min(w, h) > 0 and x + w <= size[0] and y + h <= size[1], 'Vùng ngoài canvas')
        start, length = reveal['startMs'], reveal['durationMs']
        require(isinstance(start, (int, float)) and isinstance(length, (int, float))
                and previous <= start and length >= 100 and start + length <= total_ms - 500,
                'Timing sai hoặc không giữ ảnh cuối ít nhất 0.5 giây')
        previous = start + length


def check_video(path, expected, audio=True):
    info = probe(path)
    video = [s for s in info['streams'] if s['codec_type'] == 'video']
    sound = [s for s in info['streams'] if s['codec_type'] == 'audio']
    require(video and (sound or not audio), f'Video thiếu hình/giọng đọc: {path}')
    require(abs(float(info['format']['duration']) - expected) <= 0.25, f'Sai thời lượng: {path}')
    require(abs(video[0]['width'] / video[0]['height'] - 16 / 9) < .01, 'Video không phải 16:9')
    if audio:
        require(abs(float(sound[0].get('duration', 0)) - expected) <= .25, 'Audio/video lệch thời lượng')
    # Decode the complete file to detect corrupted packets.
    command(['ffmpeg', '-v', 'error', '-xerror', '-i', path, '-f', 'null', '-'])


def stage(run, state, key, signature, outputs, action, check, attempts=3):
    record = state.setdefault('stages', {}).get(key, {})
    signature = digest(signature)
    if record.get('signature') == signature and all(p.is_file() for p in outputs):
        if record.get('outputs') == {str(p): file_hash(p) for p in outputs}:
            try:
                check()
                return
            except Exception:
                # A previously produced artifact can fail a stronger validator.
                # Rebuild it through the same bounded path.
                pass
    for attempt in range(1, attempts + 1):
        try:
            action()
            check()
            state['stages'][key] = {'signature': signature, 'outputs': {str(p): file_hash(p) for p in outputs}}
            save(run / 'state.json', state)
            return
        except Exception as exc:
            with (run / 'events.jsonl').open('a', encoding='utf-8') as stream:
                stream.write(json.dumps({'stage': key, 'attempt': attempt, 'error': str(exc)}, ensure_ascii=False) + '\n')
            if attempt == attempts:
                state['status'] = 'needs_repair'
                save(run / 'state.json', state)
                raise
            time.sleep(attempt)


def timestamp(ms):
    seconds, milliseconds = divmod(int(ms), 1000)
    minutes, seconds = divmod(seconds, 60)
    hours, minutes = divmod(minutes, 60)
    return f'{hours:02}:{minutes:02}:{seconds:02},{milliseconds:03}'


def produce(run, audio_only=False):
    state, plan = approved(run)
    board = read(run / 'storyboard.json')
    validate_storyboard(board, plan)
    require(shutil.which('ffmpeg') and shutil.which('ffprobe'), 'Cần ffmpeg và ffprobe trong PATH')
    # Engine changes invalidate cached artifacts.
    engine = digest({p.name: file_hash(p) for p in SCRIPTS.glob('*.py')})
    py = sys.executable
    clips, images, timeline, offset = [], [], [], 0
    state['status'] = 'producing'
    save(run / 'state.json', state)
    for slide in board['slides']:
        scene = run / f"scene-{slide['id']:02}"
        scene.mkdir(exist_ok=True)
        mp3, srt = scene / 'voice.mp3', scene / 'voice.srt'
        voice_text = scene / 'voice.txt'
        voice_text.write_text(slide['voice_script'], encoding='utf-8')
        def check_audio():
            require(any(s['codec_type'] == 'audio' for s in probe(mp3)['streams']), 'MP3 thiếu audio')
            validate_srt(srt, round(duration(mp3) * 1000))
        stage(run, state, f"{slide['id']}:tts", [engine, slide['voice_script'], board.get('voice')],
              [mp3, srt], lambda: command([py, SCRIPTS / 'generate_tts.py', voice_text, mp3,
                                          '--text-file', '--voice', board.get('voice', 'vi-VN-HoaiMyNeural')]), check_audio)
        if audio_only:
            continue
        from PIL import Image
        image, annotation = scene / 'slide.png', scene / 'annotation.json'
        with Image.open(image) as picture:
            size = picture.size
            require(abs(size[0] / size[1] - 16 / 9) < .01, 'Ảnh phải 16:9')
        audio_ms = math.ceil(duration(mp3) * 1000)
        # Round up to a whole video frame, retain a half-second final hold.
        total_ms = math.ceil((audio_ms + 500) * 30 / 1000) * 1000 / 30
        validate_annotation(annotation, size, total_ms)
        silent, clip = scene / 'silent.mp4', scene / 'scene.mp4'
        def render():
            command([py, SCRIPTS / 'render_stream_whiteboard.py', image, annotation, silent,
                     '--total-ms', math.ceil(total_ms), '--fps', 30, '--cap-long-edge', 1920])
            command(['ffmpeg', '-y', '-v', 'error', '-i', silent, '-i', mp3, '-map', '0:v:0',
                     '-map', '1:a:0', '-c:v', 'copy', '-c:a', 'aac', '-af', 'apad',
                     '-t', f'{total_ms / 1000:.6f}', '-movflags', '+faststart', clip])
        stage(run, state, f"{slide['id']}:render", [engine, file_hash(image), file_hash(annotation), file_hash(mp3)],
              [clip], render, lambda: check_video(clip, total_ms / 1000))
        # Use encoded clip duration for cumulative offsets, not word endpoints.
        clip_ms = round(duration(clip) * 1000)
        for cue in validate_srt(srt, audio_ms):
            timeline.append({**cue, 'startMs': cue['startMs'] + offset, 'endMs': cue['endMs'] + offset})
        offset += clip_ms
        clips.append(clip)
        images.append(image)
    if audio_only:
        state['status'] = 'audio_ready'
        save(run / 'state.json', state)
        return
    from generate_slides_pptx import create_pptx_from_images
    from pptx import Presentation
    deck = run / 'lesson.pptx'
    def check_deck():
        prs = Presentation(deck)
        require(len(prs.slides) == len(images) and all(len(s.shapes) == 1 for s in prs.slides),
                'PPTX phải có đúng một ảnh mỗi slide')
    stage(run, state, 'pptx', [engine, *[file_hash(p) for p in images]], [deck],
          lambda: create_pptx_from_images(images, deck), check_deck)
    final = run / 'lesson.mp4'
    stage(run, state, 'merge', [engine, *[file_hash(p) for p in clips]], [final],
          lambda: command([py, SCRIPTS / 'merge_scenes.py', '--inputs', *clips, '--output', final]),
          lambda: check_video(final, offset / 1000))
    subtitles = run / 'lesson.srt'
    subtitles.write_text('\n\n'.join(f"{i}\n{timestamp(c['startMs'])} --> {timestamp(c['endMs'])}\n{c['text']}"
                                         for i, c in enumerate(timeline, 1)) + '\n', encoding='utf-8')
    validate_srt(subtitles, round(duration(final) * 1000))
    state['status'] = 'awaiting_agent_qa'
    state['production_board'] = digest(board)
    state['production_inputs'] = {str(p): file_hash(p) for p in
                                  images + [p.parent / 'annotation.json' for p in images]
                                  + [p.parent / 'voice.mp3' for p in images]
                                  + [p.parent / 'voice.srt' for p in images]}
    state['deliverables'] = {str(p): file_hash(p) for p in [deck, final, subtitles]}
    save(run / 'state.json', state)
    print('Kiểm tra kỹ thuật đạt. Agent cần xem ảnh, video và nghe audio trước khi hoàn tất.')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('action', choices=['init', 'checkpoint', 'approve', 'validate', 'audio', 'produce', 'finish', 'status'])
    parser.add_argument('--run', required=True, type=Path)
    parser.add_argument('--source', default='.')
    parser.add_argument('--message', default='')
    args = parser.parse_args()
    run = args.run.resolve()
    if args.action == 'init':
        require(not (run / 'state.json').exists(), 'Đã có trạng thái; dùng status hoặc thư mục output mới')
        sources = scan_root_directory(args.source)
        save(run / 'sources.json', sources)
        require(sources['documents'] and not sources['errors'], 'Thiếu tài liệu hoặc cần xử lý lỗi trong sources.json')
        save(run / 'state.json', {'status': 'drafting', 'stages': {}})
    elif args.action == 'checkpoint':
        sources_current(run)
        plan = read(run / 'plan.json')
        validate_plan(plan)
        state = read(run / 'state.json')
        state.update(status='awaiting_approval', pending_plan=digest(plan), approved_plan=None)
        save(run / 'state.json', state)
        print(STOP)
    elif args.action == 'approve':
        sources_current(run)
        state, plan = read(run / 'state.json'), read(run / 'plan.json')
        require(unicodedata.normalize('NFC', args.message.strip()) == 'Tiếp tục', 'Chỉ nhận lệnh Tiếp tục')
        validate_plan(plan)
        require(state['status'] == 'awaiting_approval' and state['pending_plan'] == digest(plan),
                'Chưa trình bày bước 4 hoặc nội dung đã đổi; checkpoint lại')
        state.update(status='approved', approved_plan=digest(plan), approval_message=args.message)
        save(run / 'state.json', state)
    elif args.action == 'validate':
        _, plan = approved(run)
        validate_storyboard(read(run / 'storyboard.json'), plan)
        print('Kịch bản đạt kiểm tra cấu trúc.')
    elif args.action in ('audio', 'produce'):
        try:
            produce(run, audio_only=args.action == 'audio')
        except Exception as exc:
            state = read(run / 'state.json')
            if state.get('status') == 'producing':
                state['status'] = 'needs_repair'
                state['last_error'] = str(exc)
                save(run / 'state.json', state)
            raise
    elif args.action == 'finish':
        state, _ = approved(run)
        require(state['status'] == 'awaiting_agent_qa', 'Chưa đạt kiểm tra kỹ thuật')
        require(state.get('production_board') == digest(read(run / 'storyboard.json')),
                'Kịch bản đã đổi; produce và QA lại')
        require(state.get('production_inputs') == {p: file_hash(p) for p in state.get('production_inputs', {})},
                'Ảnh/annotation/audio đã đổi; produce và QA lại')
        qa = read(run / 'qa.json')
        require(qa.get('passed') is True and qa.get('findings') == [] and qa.get('evidence'),
                'QA còn lỗi hoặc thiếu bằng chứng')
        require(qa.get('deliverables') == state['deliverables']
                == {p: file_hash(p) for p in state['deliverables']}, 'QA không khớp phiên bản đầu ra')
        state['status'] = 'complete'
        save(run / 'state.json', state)
    else:
        print(json.dumps(read(run / 'state.json'), ensure_ascii=False, indent=2))


if __name__ == '__main__':
    try:
        main()
    except (ValueError, KeyError, TypeError, OSError, subprocess.CalledProcessError) as exc:
        print(f'ERROR: {exc}', file=sys.stderr)
        sys.exit(1)
