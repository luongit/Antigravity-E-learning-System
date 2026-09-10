"""
core/pptx_engine.py
Bộ đóng gói bài giảng PowerPoint (.pptx) và tài liệu PDF chuẩn tỷ lệ 16:9 Widescreen.
Mỗi slide sử dụng ảnh đồ họa chất lượng cao (1920x1080) phủ kín 100% diện tích trang.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import List

# Lọc bỏ site-packages của môi trường ảo cũ nếu xung đột phiên bản
sys.path = [p for p in sys.path if not (p.endswith(r".venv\Lib\site-packages") or p.endswith(".venv/Lib/site-packages"))]

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")


def create_pptx_from_images(image_paths: List[Path], output_pptx: Path) -> Path:
    """Tạo file PPTX 16:9 với mỗi slide là 1 ảnh phủ kín trang (13.333 x 7.5 inches)."""
    if not image_paths:
        print("⚠️ Không có ảnh nào để tạo PPTX.")
        return output_pptx

    prs = Presentation()
    # Tỷ lệ 16:9 Widescreen chuẩn quốc tế: 13.333 x 7.5 inches (33.867 x 19.05 cm)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    output_pptx.parent.mkdir(parents=True, exist_ok=True)

    for img_p in image_paths:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            str(img_p),
            left=0,
            top=0,
            width=prs.slide_width,
            height=prs.slide_height,
        )

    prs.save(str(output_pptx))
    return output_pptx


def create_pdf_from_images(image_paths: List[Path], output_pdf: Path) -> Path:
    """Tạo file PDF 16:9 chất lượng cao, mỗi trang là 1 slide 1920x1080 sắc nét."""
    if not image_paths:
        print("⚠️ Không có ảnh nào để tạo PDF.")
        return output_pdf

    output_pdf.parent.mkdir(parents=True, exist_ok=True)
    pil_images: list[Image.Image] = []

    for p in image_paths:
        img = Image.open(p).convert("RGB")
        if img.size != (1920, 1080):
            img = img.resize((1920, 1080), Image.Resampling.LANCZOS)
        pil_images.append(img)

    first_img = pil_images[0]
    rest_images = pil_images[1:] if len(pil_images) > 1 else []
    first_img.save(
        str(output_pdf),
        "PDF",
        resolution=150.0,
        save_all=True,
        append_images=rest_images,
    )
    return output_pdf


def package_slides(
    image_paths: List[Path],
    output_dir: Path,
    base_name: str = "slides",
    formats: str = "all",
) -> dict[str, Path]:
    """Đóng gói danh sách ảnh slide thành cả PPTX và/hoặc PDF."""
    results = {}
    output_dir.mkdir(parents=True, exist_ok=True)

    if formats in ("pptx", "all"):
        pptx_path = output_dir / f"{base_name}.pptx"
        create_pptx_from_images(image_paths, pptx_path)
        size_mb = pptx_path.stat().st_size / (1024 * 1024)
        print(f"✅ Đã xuất PowerPoint PPTX: {pptx_path} ({size_mb:.2f} MB)")
        results["pptx"] = pptx_path

    if formats in ("pdf", "all"):
        pdf_path = output_dir / f"{base_name}.pdf"
        create_pdf_from_images(image_paths, pdf_path)
        size_mb = pdf_path.stat().st_size / (1024 * 1024)
        print(f"✅ Đã xuất tài liệu PDF 16:9: {pdf_path} ({size_mb:.2f} MB)")
        results["pdf"] = pdf_path

    return results


def main():
    parser = argparse.ArgumentParser(description="Đóng gói Slide PowerPoint & PDF 16:9")
    parser.add_argument("source", help="Thư mục chứa ảnh slide hoặc danh sách file")
    parser.add_argument("--output-dir", default=None, help="Thư mục xuất file")
    parser.add_argument("--name", default="slides", help="Tên file đầu ra (không kèm đuôi)")
    parser.add_argument(
        "--format",
        choices=["pptx", "pdf", "all"],
        default="all",
        help="Định dạng xuất (pptx, pdf, all)",
    )
    args = parser.parse_args()

    src = Path(args.source)
    if src.is_dir():
        exts = (".png", ".jpg", ".jpeg")
        images = sorted(
            [p for p in src.iterdir() if p.suffix.lower() in exts],
            key=lambda p: int("".join(filter(str.isdigit, p.stem)) or 0),
        )
    else:
        images = [src]

    out_dir = Path(args.output_dir) if args.output_dir else (src if src.is_dir() else src.parent)
    package_slides(images, out_dir, args.name, args.format)


if __name__ == "__main__":
    main()
