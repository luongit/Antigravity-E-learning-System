"""Root-only source inventory; extraction failures never become lesson content."""
import argparse
import hashlib
import json
import sys
from pathlib import Path

SUPPORTED = {'.pdf', '.docx', '.txt', '.md', '.json'}
PROJECT_FILES = {'skill.md', 'readme.md', 'agents.md', 'requirements.txt', 'package.json', 'package-lock.json'}


def read_document(file_path):
    path = Path(file_path)
    if path.suffix.lower() == '.docx':
        from docx import Document
        doc = Document(path)
        return '\n'.join([p.text for p in doc.paragraphs] +
                         [' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
    if path.suffix.lower() == '.pdf':
        from pypdf import PdfReader
        parts = []
        for i, page in enumerate(PdfReader(path).pages, 1):
            text = page.extract_text() or ''
            if not text.strip():
                raise ValueError(f'Trang {i} không có văn bản; cần OCR/đọc ảnh')
            parts.append(f'[Trang {i}]\n{text}')
        return '\n'.join(parts)
    return path.read_text(encoding='utf-8-sig')


def scan_root_directory(root_dir='.'):
    root = Path(root_dir).resolve()
    result = {'root': str(root), 'documents': [], 'excluded': [], 'errors': []}
    for path in sorted(root.iterdir()):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED:
            continue
        if path.name.lower() in PROJECT_FILES:
            result['excluded'].append({'path': path.name, 'reason': 'Tệp cấu hình/hướng dẫn dự án'})
            continue
        try:
            content = read_document(path)
            if not content.strip():
                raise ValueError('Tài liệu rỗng hoặc cần OCR')
            result['documents'].append({'path': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'text': content})
        except Exception as exc:
            result['errors'].append({'path': path.name, 'error': str(exc)})
    return result


if __name__ == '__main__':
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    parser = argparse.ArgumentParser()
    parser.add_argument('folder', nargs='?', default='.')
    parser.add_argument('--output')
    args = parser.parse_args()
    result = scan_root_directory(args.folder)
    encoded = json.dumps(result, ensure_ascii=False, indent=2)
    if args.output:
        destination = Path(args.output)
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_text(encoded, encoding='utf-8')
    else:
        print(encoded)
    raise SystemExit(1 if result['errors'] or not result['documents'] else 0)

